"""
drive_reader.py
Fetches files from Google Drive using a Service Account and extracts their text.
Supports: PDF, Google Docs, Google Slides, .pptx, .docx, .xlsx

Changes from v1:
  - Added retry with exponential backoff on Drive API 429 / 5xx errors
  - Added .xlsx extraction support
  - Added retry helper _retry_call()
"""

import io
import os
import json
import time
import logging
from typing import Optional

from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from googleapiclient.errors import HttpError
import pdfplumber
from pptx import Presentation
import docx

logger = logging.getLogger(__name__)

# ── MIME TYPE ROUTING ─────────────────────────────────────────────────────────

EXTRACTABLE_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.google-apps.document": "gdoc",
    "application/vnd.google-apps.presentation": "gslides",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    # NEW: Excel support
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel": "xlsx",
    "application/vnd.google-apps.spreadsheet": "gsheet",
}

SKIP_TYPES = {
    "video/mp4",
    "video/quicktime",
    "video/mp2t",
    "video/x-m4v",
    "audio/mpeg",
    "audio/mp4",
    "audio/x-m4a",
    "image/png",
    "image/jpeg",
    "image/gif",
    "image/webp",
    "application/octet-stream",
}

# Upper bound on text kept per document. This used to be 6,000 characters,
# which meant only the first 2–3 pages of every manual were searchable.
MAX_CHARS_PER_DOC = int(os.environ.get("MAX_CHARS_PER_DOC", "400000"))
MAX_RETRIES = 3
RETRY_BACKOFF = 2  # seconds; doubles each retry


# ── SERVICE ACCOUNT SETUP ─────────────────────────────────────────────────────

def get_drive_service():
    scopes = ["https://www.googleapis.com/auth/drive.readonly"]

    creds_raw = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
    if not creds_raw:
        raise EnvironmentError(
            "GOOGLE_SERVICE_ACCOUNT_JSON environment variable is not set. "
            "See .env.example for instructions."
        )

    try:
        creds_dict = json.loads(creds_raw)
    except json.JSONDecodeError as e:
        raise ValueError(f"GOOGLE_SERVICE_ACCOUNT_JSON is not valid JSON: {e}")

    creds = service_account.Credentials.from_service_account_info(
        creds_dict, scopes=scopes
    )
    return build("drive", "v3", credentials=creds, cache_discovery=False)


# ── FILE ID EXTRACTION ────────────────────────────────────────────────────────

def extract_file_id(drive_url: str) -> Optional[str]:
    import re

    patterns = [
        r"/file/d/([a-zA-Z0-9_-]+)",
        r"/document/d/([a-zA-Z0-9_-]+)",
        r"/presentation/d/([a-zA-Z0-9_-]+)",
        r"/spreadsheets/d/([a-zA-Z0-9_-]+)",
        r"[?&]id=([a-zA-Z0-9_-]+)",
        r"/open\?id=([a-zA-Z0-9_-]+)",
    ]

    for pattern in patterns:
        match = re.search(pattern, drive_url)
        if match:
            return match.group(1)

    logger.warning(f"Could not extract file ID from URL: {drive_url}")
    return None


# ── RETRY HELPER ─────────────────────────────────────────────────────────────

def _retry_call(fn, *args, max_retries=MAX_RETRIES, **kwargs):
    """
    Calls fn(*args, **kwargs) and retries on HttpError 429/5xx with
    exponential backoff. Raises on final failure.
    """
    delay = RETRY_BACKOFF
    for attempt in range(1, max_retries + 1):
        try:
            return fn(*args, **kwargs)
        except HttpError as e:
            status = e.resp.status if hasattr(e, "resp") else 0
            if status in (429, 500, 502, 503, 504) and attempt < max_retries:
                logger.warning(f"Drive API error {status}, retry {attempt}/{max_retries} in {delay}s...")
                time.sleep(delay)
                delay *= 2
            else:
                raise


# ── MAIN ENTRY POINT ──────────────────────────────────────────────────────────

def is_media(mime_type: str) -> bool:
    return mime_type in SKIP_TYPES or mime_type.split("/")[0] in ("video", "audio", "image")


def is_extractable(mime_type: str) -> bool:
    return mime_type in EXTRACTABLE_TYPES


def extract_text(file_id: str, mime_type: str, name: str, service) -> str:
    """
    Extracts text from a Drive file whose metadata is already known.
    Raises on failure so the caller can record the error.
    """
    file_type = EXTRACTABLE_TYPES[mime_type]
    extractor = {
        "pdf": _extract_pdf,
        "gdoc": _extract_google_doc,
        "gslides": _extract_google_slides,
        "gsheet": _extract_google_sheet,
        "pptx": _extract_pptx,
        "docx": _extract_docx,
        "xlsx": _extract_xlsx,
    }[file_type]
    return (extractor(file_id, service, name) or "").replace("\x00", "")


def extract_text_from_drive_file(drive_url: str, service) -> Optional[str]:
    file_id = extract_file_id(drive_url)
    if not file_id:
        return None

    try:
        metadata = _retry_call(
            service.files().get(
                fileId=file_id,
                fields="id, name, mimeType",
                supportsAllDrives=True,
            ).execute
        )
    except Exception as e:
        logger.error(f"Failed to get metadata for file_id={file_id}: {e}")
        return None

    mime_type = metadata.get("mimeType", "")
    name = metadata.get("name", "unknown")

    logger.info(f"Processing: {name} | MIME: {mime_type}")

    if not is_extractable(mime_type):
        logger.info(f"  → Skipping (media or unsupported type {mime_type}): {name}")
        return None

    try:
        return extract_text(file_id, mime_type, name, service)
    except Exception as e:
        logger.error(f"  → Extraction failed for {name}: {e}")
        return None


# ── EXTRACTORS ────────────────────────────────────────────────────────────────

def _download_to_buffer(file_id: str, service) -> io.BytesIO:
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    buffer = io.BytesIO()
    downloader = MediaIoBaseDownload(buffer, request)

    done = False
    while not done:
        _, done = downloader.next_chunk()

    buffer.seek(0)
    return buffer


def _extract_pdf(file_id: str, service, name: str) -> str:
    logger.info(f"  → Extracting PDF: {name}")
    buffer = _download_to_buffer(file_id, service)
    text = ""

    with pdfplumber.open(buffer) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
            # Free each page's parsed objects; big manuals otherwise use a
            # lot of memory on a small server.
            if hasattr(page, "close"):
                page.close()
            if len(text) >= MAX_CHARS_PER_DOC:
                break

    return text[:MAX_CHARS_PER_DOC]


def _extract_google_doc(file_id: str, service, name: str) -> str:
    logger.info(f"  → Exporting Google Doc: {name}")
    content = _retry_call(
        service.files().export(fileId=file_id, mimeType="text/plain").execute
    )
    return content.decode("utf-8")[:MAX_CHARS_PER_DOC]


def _extract_google_slides(file_id: str, service, name: str) -> str:
    logger.info(f"  → Exporting Google Slides: {name}")
    content = _retry_call(
        service.files().export(fileId=file_id, mimeType="text/plain").execute
    )
    return content.decode("utf-8")[:MAX_CHARS_PER_DOC]


def _extract_google_sheet(file_id: str, service, name: str) -> str:
    """Exports a Google Sheet as CSV text."""
    logger.info(f"  → Exporting Google Sheet: {name}")
    content = _retry_call(
        service.files().export(fileId=file_id, mimeType="text/csv").execute
    )
    return content.decode("utf-8")[:MAX_CHARS_PER_DOC]


def _extract_pptx(file_id: str, service, name: str) -> str:
    logger.info(f"  → Extracting PPTX: {name}")
    buffer = _download_to_buffer(file_id, service)
    prs = Presentation(buffer)
    text = ""

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text += shape.text.strip() + "\n"
        if len(text) >= MAX_CHARS_PER_DOC:
            break

    return text[:MAX_CHARS_PER_DOC]


def _extract_docx(file_id: str, service, name: str) -> str:
    logger.info(f"  → Extracting DOCX: {name}")
    buffer = _download_to_buffer(file_id, service)
    doc = docx.Document(buffer)
    text = "\n".join(
        para.text for para in doc.paragraphs if para.text.strip()
    )
    return text[:MAX_CHARS_PER_DOC]


def _extract_xlsx(file_id: str, service, name: str) -> str:
    """Downloads and extracts text from a .xlsx file."""
    logger.info(f"  → Extracting XLSX: {name}")
    import openpyxl
    buffer = _download_to_buffer(file_id, service)
    wb = openpyxl.load_workbook(buffer, read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(cell) if cell is not None else "" for cell in row)
            if row_text.strip():
                lines.append(row_text)
        if sum(len(l) for l in lines) >= MAX_CHARS_PER_DOC:
            break
    return "\n".join(lines)[:MAX_CHARS_PER_DOC]
